import os
import re
from openpyxl import load_workbook
from openpyxl_image_loader import SheetImageLoader

# ============ 配置区 ============
XLSX_PATH = r"MTK9655-LF1V107.xlsx"           # 表格路径
OUTPUT_DIR = r"./template"               # 输出根目录
SHEET_NAME = None                      # 指定工作表名，None 表示遍历所有
# ================================


def sanitize(name):
    """清理文件名中的非法字符"""
    return re.sub(r'[\\/:*?"<>|]', '_', str(name).strip())


def extract_sid(sid_cell_value):
    """从 'sid = "0"' 中提取 0"""
    if sid_cell_value is None:
        return None
    text = str(sid_cell_value)
    match = re.search(r'(\d+)', text)
    return match.group(1) if match else None


def save_image(image_loader, cell_coordinate, save_path):
    """安全保存单元格中的图片"""
    try:
        if image_loader.image_in(cell_coordinate):
            img = image_loader.get(cell_coordinate)
            # 统一转为 PNG 保存
            img.save(save_path, format="PNG")
            return True
    except Exception as e:
        print(f"  ⚠️ 保存图片失败 {cell_coordinate}: {e}")
    return False


def main():
    wb = load_workbook(XLSX_PATH)

    sheets = [SHEET_NAME] if SHEET_NAME else wb.sheetnames

    for sheet_name in sheets:
        sheet = wb[sheet_name]
        print(f"\n📄 处理工作表：{sheet_name}")
        image_loader = SheetImageLoader(sheet)

        # 从第 2 行开始（假设第 1 行是表头）
        for row in range(2, sheet.max_row + 1):
            sid_raw = sheet.cell(row, 1).value   # A列：SID
            model   = sheet.cell(row, 2).value   # B列：机型

            if sid_raw is None or model is None:
                continue

            sid_num = extract_sid(sid_raw)
            if sid_num is None:
                print(f"  第{row}行 SID 解析失败：{sid_raw}")
                continue

            model = sanitize(model)
            folder_name = f"SID_{sid_num}_{model}"
            folder_path = os.path.join(OUTPUT_DIR, folder_name)
            os.makedirs(folder_path, exist_ok=True)

            print(f"  📁 {folder_name}")

            # C列：开机logo → 开机logo.png
            logo_path = os.path.join(folder_path, "开机logo.png")
            if save_image(image_loader, f"C{row}", logo_path):
                print(f"    ✅ 开机logo 已保存")
            else:
                print(f"    ⚠️ 第{row}行 C列无图片")

            # D列：开机动画 → 开机动画.png
            anim_path = os.path.join(folder_path, "开机动画.png")
            if save_image(image_loader, f"D{row}", anim_path):
                print(f"    ✅ 开机动画 已保存")
            else:
                print(f"    ⚠️ 第{row}行 D列无图片")

    print("\n🎉 全部完成！输出目录：", os.path.abspath(OUTPUT_DIR))


if __name__ == "__main__":
    main()